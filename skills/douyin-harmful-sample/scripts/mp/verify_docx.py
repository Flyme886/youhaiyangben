#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""结构校验（MINIMAL_LAYOUT_V1）：每份 ≤50 条、编号连续、标题→链接成对、无旧版式残留。"""
import argparse
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

FORBIDDEN = ["【违规要素】", "【概括】", "原文：", "链接：", "本批编号范围", "举报理由草稿", "待补", "人工核验"]


def check(path: Path, expect_start: int):
    errs = []
    doc = Document(str(path))
    if len(doc.tables):
        errs.append(f"{path.name}: 含表格 {len(doc.tables)} 个")
    paras = doc.paragraphs
    if len(paras) % 3 != 0:
        errs.append(f"{path.name}: 段落数 {len(paras)} 不是 3 的倍数")
    n_items = len(paras) // 3
    if n_items > 50:
        errs.append(f"{path.name}: 条数 {n_items} > 50")
    body = "\n".join(p.text for p in paras)
    for bad in FORBIDDEN:
        if bad in body:
            errs.append(f"{path.name}: 残留禁用文案「{bad}」")

    expected = expect_start
    for i in range(n_items):
        tp, lp, ip = paras[i * 3], paras[i * 3 + 1], paras[i * 3 + 2]
        t = tp.text.strip()
        if not t.startswith(f"{expected}."):
            errs.append(f"{path.name}: 第 {i+1} 组编号异常，期望 {expected}. 实得「{t[:24]}」")
        expected += 1
        if tp.runs:
            r = tp.runs[0]
            if not r.bold:
                errs.append(f"{path.name}: 条 {i+1} 标题未加粗")
            if r.font.size and r.font.size != Pt(12):
                errs.append(f"{path.name}: 条 {i+1} 标题字号 {r.font.size}")
            if r.font.color and r.font.color.rgb is not None and str(r.font.color.rgb) != "1F4E79":
                errs.append(f"{path.name}: 条 {i+1} 标题颜色 {r.font.color.rgb}")
        else:
            errs.append(f"{path.name}: 条 {i+1} 标题段无 run")
        if not lp.text.strip().startswith("http"):
            errs.append(f"{path.name}: 条 {i+1} 链接段非 http 开头")
        if lp.runs and lp.runs[0].font.size and lp.runs[0].font.size != Pt(9.5):
            errs.append(f"{path.name}: 条 {i+1} 链接字号 {lp.runs[0].font.size}")
        if ip.text.strip():
            errs.append(f"{path.name}: 条 {i+1} 图片段有文字「{ip.text[:20]}」")

    shapes = doc.inline_shapes
    for i in range(len(shapes)):
        sh = shapes[i]
        if sh.width.cm > 16.05:
            errs.append(f"{path.name}: 图 {i+1} 宽 {sh.width.cm:.2f}cm 超限")
        if sh.height.cm > 20.55:
            errs.append(f"{path.name}: 图 {i+1} 高 {sh.height.cm:.2f}cm 超限")
    return n_items, errs, len(shapes)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("dir")
    args = ap.parse_args()
    d = Path(args.dir)
    files = sorted(d.glob("*.docx"))
    if not files:
        print("未找到 docx")
        return 1
    total = imgs = 0
    errs_all = []
    for f in files:
        n, errs, im = check(f, total + 1)
        total += n
        imgs += im
        errs_all += errs
        print(f"{f.name}: {n} 条, 图 {im} 张")
    print(f"合计 {total} 条 / 图 {imgs} 张 / 缺图 {total - imgs} 条")
    if errs_all:
        print("=== 失败 ===")
        for e in errs_all[:40]:
            print(" -", e)
        return 2
    print("=== 结构校验通过 ===")
    return 0


if __name__ == "__main__":
    sys.exit(main())
