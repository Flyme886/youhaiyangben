#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""MINIMAL_LAYOUT_V1 Word 生成器（参数化，可多平台复用）

每条样本严格三段：N.标题 / 原始链接 / 页面截图
- 标题 12pt 加粗 #1F4E79，段前 2pt
- 链接 9.5pt #0B60B0，段前 4pt
- 截图段居中，横图固定 16.0cm 宽，竖图按高 20.5cm 限幅
- 截图缺失则该位置留空白段落，不写任何说明文字
"""
import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as DocxImage
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

COLOR_TITLE = RGBColor(0x1F, 0x4E, 0x79)
COLOR_LINK = RGBColor(0x0B, 0x60, 0xB0)
MAX_W, MAX_H = 16.0, 20.5
DELETED_HINT = ("不存在", "已删除", "已被删除", "无法查看", "帖子可能已被删除")


def usable(shot: dict | None) -> bool:
    """ok = 抓到正文；deleted = 页面明说已删除，同样是有效证据。"""
    if not shot:
        return False
    if shot.get("status") in ("ok", "deleted"):
        return True
    text = (shot.get("page_text") or "") + " " + (shot.get("page_title") or "")
    return any(h in text for h in DELETED_HINT)


def load_shots(dirs: list[Path]) -> dict:
    out = {}
    for d in dirs:
        if not d.exists():
            continue
        for f in sorted(d.glob("results_*.jsonl")):
            for line in f.read_text(encoding="utf-8").splitlines():
                if not line.strip():
                    continue
                try:
                    o = json.loads(line)
                except json.JSONDecodeError:
                    continue
                k = str(o.get("shot_key"))
                if k not in out or o.get("status") == "ok":
                    out[k] = o
    return out


def new_doc() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        fonts.set(qn(attr), "微软雅黑")
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(3.175)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    return doc


def add_item(doc, n: int, item: dict, shot: dict | None, shots_dir: Path) -> bool:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(f"{n}.{item['title']}")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_TITLE

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(4)
    r2 = p2.add_run(item["link"])
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = COLOR_LINK

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img = shots_dir / f"{item['shot_key']}.png"
    if shot and usable(shot) and img.exists():
        try:
            info = DocxImage.from_file(str(img))
            w_px, h_px = info.px_width, info.px_height
            if h_px > w_px:
                p3.add_run().add_picture(
                    str(img), width=Cm(min(MAX_W, MAX_H * w_px / h_px)), height=Cm(MAX_H)
                )
            else:
                p3.add_run().add_picture(str(img), width=Cm(MAX_W))
            return True
        except Exception as e:
            print(f"  ! 插图失败 {img.name}: {e}")
    return False


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--manifest", required=True)
    ap.add_argument("--shots", required=True)
    ap.add_argument("--results", nargs="+", default=[])
    ap.add_argument("--outdir", required=True)
    ap.add_argument("--prefix", default="第")
    ap.add_argument("--batch", type=int, default=50)
    args = ap.parse_args()

    items = json.loads(Path(args.manifest).read_text(encoding="utf-8"))["items"]
    shots_dir = Path(args.shots)
    shots = load_shots([Path(p) for p in args.results])
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    nb = (len(items) + args.batch - 1) // args.batch
    missing = []
    for b in range(nb):
        chunk = items[b * args.batch : (b + 1) * args.batch]
        doc = new_doc()
        got = 0
        for idx, item in enumerate(chunk, start=b * args.batch + 1):
            shot = shots.get(str(item["shot_key"]))
            if add_item(doc, idx, item, shot, shots_dir):
                got += 1
            else:
                missing.append((idx, item.get("platform_name") or item.get("platform"), item.get("author", "")))
        name = f"{args.prefix}{b + 1:02d}批.docx"
        doc.save(outdir / name)
        print(f"{name}: {len(chunk)} 条，附图 {got}")

    print(f"合计 {len(items)} 条 / 缺图 {len(missing)} 条")
    for idx, plat, author in missing:
        print(f"  缺图 {idx}. [{plat}] {author}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
